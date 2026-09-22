# file_reader.py
# v2: DOCX tables in document order, PDF page markers for RAG,
#     no API changes.

import os
import logging

logger = logging.getLogger(__name__)


def read_txt(path):
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            with open(path, "r", encoding=enc) as f:
                return f.read()
        except Exception:
            continue
    return ""


def read_pdf(path):
    """
    Extract text page by page with [PAGE N] markers so downstream
    code (RAG, checker) can cite page numbers.
    """
    try:
        import fitz
        doc = fitz.open(path)
        parts = []
        for i, page in enumerate(doc, start=1):
            text = page.get_text("text")
            if text.strip():
                parts.append(f"[PAGE {i}]\n{text}")
        doc.close()
        return "\n\n".join(parts)
    except Exception as e:
        logger.error("PDF read failed: %s", e)
        return ""


def read_xlsx(path):
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True, read_only=True)
        parts = []
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            parts.append(f"# Sheet: {sheet}")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                if any(cells):
                    parts.append(" | ".join(cells))
        wb.close()
        return "\n".join(parts)
    except Exception as e:
        logger.error("XLSX read failed: %s", e)
        return ""


def read_docx(path):
    """
    Extract paragraphs AND tables from a DOCX, preserving order.
    Egyptian site documents are table-heavy; skipping tables loses
    most of the content.
    """
    try:
        from docx import Document
        from docx.oxml.ns import qn
        doc = Document(path)
        out = []
        body = doc.element.body
        for child in body.iterchildren():
            tag = child.tag
            if tag == qn("w:p"):
                para_text = "".join(
                    node.text or "" for node in child.iter()
                    if node.tag == qn("w:t"))
                if para_text.strip():
                    out.append(para_text)
            elif tag == qn("w:tbl"):
                for row in child.iter(qn("w:tr")):
                    cells = []
                    for cell in row.iter(qn("w:tc")):
                        cell_text = " ".join(
                            node.text or "" for node in cell.iter()
                            if node.tag == qn("w:t"))
                        cells.append(cell_text.strip())
                    if any(cells):
                        out.append(" | ".join(cells))
                out.append("")
        return "\n".join(out)
    except Exception as e:
        logger.error("DOCX read failed: %s", e)
        return ""


async def read_image(path):
    from gemini import ocr_image
    return await ocr_image(path)


async def extract_text(path, filename):
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".pdf":
        return read_pdf(path), "pdf"
    if ext == ".txt":
        return read_txt(path), "txt"
    if ext in (".xlsx", ".xlsm"):
        return read_xlsx(path), "xlsx"
    if ext in (".png", ".jpg", ".jpeg", ".webp", ".bmp"):
        return await read_image(path), "image"
    if ext == ".docx":
        return read_docx(path), "docx"
    return "", "unknown"
